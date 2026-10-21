"""Builds the DigitVision project presentation (9 slides).
ACC is filled in from the trained weights.json so the results slide
shows the real 5-fold CV number.
"""
import json

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls


def set_cell_borders(cell, color="000000", w="12700"):  # w in EMU; 12700 = 1pt
    """Give a table cell plain black grid borders via raw XML
    (python-pptx exposes no public border API)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnT", "a:lnL", "a:lnB", "a:lnR"):
        ln = parse_xml(
            f'<{tag} {nsdecls("a")} w="{w}">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill></{tag}>'
        )
        tcPr.append(ln)

with open("weights.json") as f:
    W = json.load(f)
ACC = W["accuracy"]
ACC_PCT = f"{ACC * 100:.1f}%"
N = W["n_samples"]
FOLDS = W.get("fold_accuracies", [])

BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x44, 0x44, 0x44)
FONT = "Times New Roman"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def txbox(slide, left, top, width, height, text, size=20, bold=False,
          color=BLACK, align=None, font=FONT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    if align is not None:
        p.alignment = align
    return tb


def number_tag(slide, n):
    txbox(slide, 0.8, 0.35, 2, 0.6, f"{n:02d}", size=18, bold=True, color=GREY)


def content_slide(title, n, body_lines, table_rows=None):
    """Title + body text lines, optional plain black grid table under the text."""
    s = prs.slides.add_slide(blank)
    number_tag(s, n)
    txbox(s, 0.8, 0.85, 11.7, 1.1, title, size=36, bold=True)
    y = 2.05
    for line in body_lines:
        txbox(s, 0.8, y, 11.7, 0.6, line, size=20)
        y += 0.55
    if table_rows:
        rows, cols = len(table_rows), len(table_rows[0])
        g = s.shapes.add_table(rows, cols, Inches(0.8), Inches(y + 0.25),
                               Inches(8.5), Inches(0.45 * rows))
        tbl = g.table
        for r in range(rows):
            for c in range(cols):
                cell = tbl.cell(r, c)
                cell.text = table_rows[r][c]
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(16)
                    p.font.name = FONT
                    p.font.color.rgb = BLACK
                    p.font.bold = (r == 0)
                set_cell_borders(cell, "000000", "12700")
        for r in range(rows):
            tbl.rows[r].height = Inches(0.45)
    return s


# 1 — Title
s = prs.slides.add_slide(blank)
txbox(s, 1.0, 1.6, 11.3, 1.4, "DigitVision", size=56, bold=True, align=PP_ALIGN.CENTER)
txbox(s, 1.0, 3.1, 11.3, 1.0, "AI Handwritten Digit Classifier", size=28,
      align=PP_ALIGN.CENTER)
txbox(s, 1.0, 4.4, 11.3, 2.0,
      "Vikranth Reddy  |  B.Tech CSE (AI & ML), Woxsen University\n"
      "Daily AI/ML Project Factory — Day 4  ·  21 October 2026",
      size=20, color=GREY, align=PP_ALIGN.CENTER)

# 2 — Problem
content_slide("The Problem", 2, [
    "•  Handwriting is messy: every person draws digits differently — slanted,",
    "    tilted, thin or thick strokes.",
    "•  Rule-based templates fail: a fixed shape for '7' never matches real writing.",
    "•  Banks, exam sheets and postal services scan millions of handwritten digits —",
    "    they need a model that learns what each digit looks like from examples.",
    "•  Goal: draw a digit 0–9 → get an instant prediction from a real ML model.",
])

# 3 — Solution
content_slide("The Solution", 3, [
    "•  DigitVision: a handwritten digit classifier trained on 1,797 8×8 grayscale",
    "    digit images (classes 0–9).",
    "•  Pipeline: drawn image → downsample to 8×8 → 64 pixel values →",
    "    softmax regression → per-digit probabilities.",
    "•  Trained weights (10 weight maps + 10 biases) are exported to weights.json —",
    "    the website re-runs the exact same math in JavaScript. Zero backend.",
    "•  Verified: browser predictions match sklearn's predict_proba exactly.",
])

# 4 — Tech Stack
content_slide("Tech Stack", 4, [
    "•  Python, scikit-learn, NumPy — training & evaluation (train.py).",
    "•  Softmax regression (multinomial logistic regression, LBFGS solver): one",
    "    64-pixel weight vector + bias per digit class; pixels scaled to 0–1.",
    "•  Stratified 5-fold cross-validation for an honest accuracy number.",
    "•  Vanilla JavaScript + HTML5 canvas — draw, downsample, softmax, probability bars.",
    "•  Static hosting on Vercel with hardened security headers.",
])

# 5 — Architecture
content_slide("Architecture", 5, [
    "•  train.py: (1) load sklearn digits dataset (1,797 images) → (2) scale pixels",
    "    to 0–1 → (3) stratified 5-fold CV → (4) fit softmax on all data →",
    "    (5) export coef_, intercept_ and 10 labeled sample digits to weights.json.",
    "•  weights.json: 10 × 64 weights + 10 biases — tiny enough to ship to the browser.",
    "•  index.html: pointer drawing → 280×280 canvas downsampled to 8×8 →",
    "    linear scores + softmax → prediction + 0–9 probability bars + 8×8 preview.",
    "•  Same input in Python and JS gives the same probability — parity verified.",
])

# 6 — Demo
content_slide("Live Demo", 6, [
    "•  Draw any digit 0–9 in the black canvas, hit Classify → prediction with a",
    "    confidence bar for every digit 0–9.",
    "•  The 8×8 preview shows exactly what the model sees after downsampling.",
    "•  One-click buttons load a real labeled sample digit from the training set.",
    "•  Demo URL: digitvision-hvrgaming2705s-projects.vercel.app",
])

# 7 — Results (with plain black grid table of fold scores)
fold_pct = [f"{a * 100:.1f}" for a in FOLDS] if FOLDS else ["—"] * 5
fold_rows = [["Fold", "1", "2", "3", "4", "5", "Mean"],
             ["Accuracy (%)"] + fold_pct + [f"{ACC * 100:.1f}"]]
content_slide("Results", 7, [
    f"•  5-fold cross-validation accuracy: {ACC_PCT} ({N} images, classes 0–9, stratified folds).",
    "•  Model: softmax regression on 64 normalized pixel features.",
    "•  Learned weight maps visibly look like the digits they score — the model",
    "    is a weighted vote of pixels, fully interpretable.",
    "•  Known limitation: trained on clean 8×8 images, so very stylized writing",
    "    (e.g. European 7 with a bar) can confuse it — a CNN would generalize better.",
], table_rows=fold_rows)

# 8 — Future Scope
content_slide("Future Scope", 8, [
    "•  Train a small convolutional neural network (CNN) — state-of-the-art for digits.",
    "•  Scale up to MNIST (70,000 images) and then EMNIST letters.",
    "•  Data augmentation: rotations, shifts, elastic distortions for robustness.",
    "•  Multi-digit sequence reading (e.g. read a whole handwritten PIN).",
    "•  Deploy the trained CNN with TensorFlow.js for in-browser inference.",
])

# 9 — Thank You
s = prs.slides.add_slide(blank)
txbox(s, 1.5, 2.2, 10.3, 1.6, "Thank You", size=60, bold=True, align=PP_ALIGN.CENTER)
txbox(s, 1.5, 4.3, 10.3, 1.5,
      "GitHub: github.com/HVRGAMING2705/digitvision\nDemo: digitvision-hvrgaming2705s-projects.vercel.app",
      size=22, color=GREY, align=PP_ALIGN.CENTER)

out = "/home/hatch/workspace/goals/daily-ai-ml-project-factory/files/AIML_DigitVision_2026-10-21.pptx"
prs.save(out)
print("saved", out)
print(f"slides: {len(prs.slides)}, CV accuracy baked in: {ACC_PCT}")
